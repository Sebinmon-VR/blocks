from docx import Document

doc = Document()
doc.add_heading("Sample DOCX File", level=1)
doc.add_paragraph("This is a short Word file with some text.")
doc.add_paragraph("Second line of sample content.")
doc.save("samples/sample.docx")
print("DOCX created")
